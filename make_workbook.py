"""Build AHGSL_results.xlsx — the full experimental record backing the manuscript."""
import json
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

agg = json.load(open("results/agg.json"))
YE = json.load(open("results/yeast_results.json"))
TC = json.load(open("results/tcga_results.json"))
GS = json.load(open("results/tcga_gene_selection.json"))

HDR = Font(name="Arial", bold=True, size=10)
TXT = Font(name="Arial", size=10)
TITLE = Font(name="Arial", bold=True, size=12)
FILL = PatternFill("solid", fgColor="DCE6F1")
thin = Side(style="thin", color="808080")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
DATASETS = ["cellcycle", "church", "derisi", "eisen", "yeast14"]
YE_M = [("br_lr", "BR-LR"), ("mlp", "MLP"), ("gcn", "GCN"),
        ("static_hgnn", "Static-HGNN"), ("ahgsl", "AHGSL (proposed)")]
ABL = [("ahgsl", "AHGSL (full)"), ("ahgsl_no_aee", "w/o AEE"),
       ("ahgsl_no_ser", "w/o SER"), ("ahgsl_no_gate", "w/o gene gate"),
       ("ahgsl_no_stab", "w/o stability")]
TC_M = [("lr", "Softmax-LR"), ("mlp", "MLP"), ("gcn", "GCN"),
        ("static_hgnn", "Static-HGNN"), ("ahgsl", "AHGSL (proposed)")]
YM_METR = [("accuracy_ex", "Accuracy (ex.)"), ("precision_micro", "Precision (mi.)"),
           ("recall_micro", "Recall (mi.)"), ("f1_micro", "F1 (mi.)"),
           ("f1_macro", "F1 (ma.)"), ("auc_micro", "AUC (mi.)"),
           ("avg_precision_micro", "AP (mi.)"), ("ranking_loss", "Ranking loss"),
           ("hamming", "Hamming loss")]

wb = Workbook()


def sheet(name):
    ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = name
    return ws


def put_row(ws, r, vals, font=TXT, fill=None, border=True):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.font = font
        cell.alignment = Alignment(horizontal="center" if c > 1 else "left",
                                   vertical="center")
        if fill:
            cell.fill = fill
        if border:
            cell.border = BORD


def autow(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i) if i < 26 else 'A' + chr(39 + i)].width = w


# ---- Summary ----
ws = sheet("Summary")
ws.cell(1, 1, "AHGSL — Experimental Results Summary").font = TITLE
rows = [
    ("Generated", "Automated experimental pipeline (PyTorch implementation of the 13-equation AHGSL specification)"),
    ("Multi-label benchmarks", "Yeast FunCat: cellcycle, church, derisi, eisen (predefined splits; top-level FunCat labels) + canonical Mulan yeast-14 (Elisseeff & Weston; ready-made 14-label matrix, official train/test split)"),
    ("Multiclass benchmark", "TCGA-PANCAN-HiSeq-801x20531 (5 tumor types; stratified 5-fold CV; 3,000 top-variance gene nodes per fold)"),
    ("Models", "AHGSL (full) vs Static-HGNN, GCN, MLP, BR-LR/Softmax-LR + 4 ablations"),
    ("Seeds", "Yeast: 3 seeds (42, 1, 7) main models, 2 seeds (42, 1) ablations; TCGA: 5 folds, seed 42"),
]
r = 3
for k, v in rows:
    ws.cell(r, 1, k).font = HDR
    ws.cell(r, 2, v).font = TXT
    r += 1
autow(ws, [26, 130])

# ---- Yeast main ----
ws = sheet("Yeast main results")
r = 1
for d in DATASETS:
    ws.cell(r, 1, d).font = TITLE
    r += 1
    put_row(ws, r, ["Method"] + [lbl for _, lbl in YM_METR], HDR, FILL)
    r += 1
    for key, name in YE_M:
        v = agg["yeast"][d][key]
        row = [name]
        for mk, _ in YM_METR:
            m, s, n = v.get(mk, (np.nan, np.nan, 0))
            row.append(f"{m:.4f} ± {s:.4f}" if n > 1 else f"{m:.4f}")
        put_row(ws, r, row)
        r += 1
    r += 1
autow(ws, [20] + [15] * 9)

# ---- Yeast ablation ----
ws = sheet("Yeast ablation")
r = 1
for d in DATASETS:
    ws.cell(r, 1, d).font = TITLE
    r += 1
    put_row(ws, r, ["Configuration", "F1 (mi.)", "AUC (mi.)", "AP (mi.)",
                    "Refined frac (mean)", "Epoch time (s)"], HDR, FILL)
    r += 1
    for key, name in ABL:
        v = agg["yeast_abl"][d][key]
        def g(mk, prec=4):
            m, s, n = v.get(mk, (np.nan, np.nan, 0))
            return f"{m:.{prec}f} ± {s:.{prec}f}" if n > 1 else f"{m:.{prec}f}"
        put_row(ws, r, [name, g("f1_micro"), g("auc_micro"),
                        g("avg_precision_micro"), g("mean_refined_frac", 3),
                        g("mean_epoch_time", 2)])
        r += 1
    r += 1
autow(ws, [22, 18, 18, 18, 18, 16])

# ---- TCGA ----
ws = sheet("TCGA classification")
put_row(ws, 1, ["Method", "Accuracy", "Precision (ma.)", "Recall (ma.)",
                "F1 (ma.)", "AUC (OvR)", "Train time (s)"], HDR, FILL)
r = 2
for key, name in TC_M + [(k, n) for k, n in ABL[1:]]:
    v = agg["tcga"][key]
    def g(mk, prec=4):
        m, s, n = v.get(mk, (np.nan, np.nan, 0))
        return f"{m:.{prec}f} ± {s:.{prec}f}" if n > 1 else f"{m:.{prec}f}"
    put_row(ws, r, [name, g("accuracy"), g("precision_macro"), g("recall_macro"),
                    g("f1_macro"), g("auc_macro_ovr"), g("train_time", 1)])
    r += 1
r += 1
ws.cell(r, 1, "Pooled confusion matrix (AHGSL, 5 folds)").font = HDR
r += 1
classes = ["BRCA", "COAD", "KIRC", "LUAD", "PRAD"]
put_row(ws, r, ["True \\ Pred"] + classes, HDR, FILL)
r += 1
for i, row in enumerate(agg["tcga_confusion"]):
    put_row(ws, r, [classes[i]] + row)
    r += 1
autow(ws, [22, 18, 18, 18, 18, 18, 14])

# ---- Gene selection ----
ws = sheet("TCGA gene selection")
put_row(ws, 1, ["Ranking method", "K=25", "K=50", "K=100", "K=200", "K=500",
                "Stability (Jaccard top-100)"], HDR, FILL)
import itertools
r = 2
for meth in ["AHGSL gate", "ANOVA F-test", "Mutual information", "L1-logistic", "Random"]:
    row = [meth]
    for K in (25, 50, 100, 200, 500):
        row.append(round(100 * np.mean([GS[f]["curves"][meth][str(K)]["accuracy"]
                                        for f in GS]), 2))
    tops = [set(GS[f]["gene_names"][i] for i in GS[f]["top100"][meth]) for f in GS]
    jac = np.mean([len(a & b) / len(a | b) for a, b in itertools.combinations(tops, 2)])
    row.append(round(float(jac), 3))
    put_row(ws, r, row)
    r += 1
r += 1
ws.cell(r, 1, "Top-50 gate-ranked genes per fold (UCI anonymized identifiers)").font = HDR
r += 1
for f in sorted(GS):
    names = GS[f]["gene_names"]
    top = [names[i] for i in GS[f]["top100"]["AHGSL gate"][:50]]
    put_row(ws, r, [f"fold {int(f[1])+1}"] + top[:25], border=False)
    put_row(ws, r + 1, [""] + top[25:], border=False)
    r += 2
autow(ws, [22] + [11] * 26)

# ---- Per-run raw ----
ws = sheet("Per-run raw (yeast)")
metr_keys = [mk for mk, _ in YM_METR] + ["train_time", "epochs_run", "thr"]
put_row(ws, 1, ["Dataset", "Method", "Seed"] + metr_keys, HDR, FILL)
r = 2
for key in sorted(YE):
    d, meth, seed = key.split("|")
    v = YE[key]
    if "error" in v:
        continue
    put_row(ws, r, [d, meth, seed] + [round(v[mk], 5) if mk in v else "" for mk in metr_keys])
    r += 1
autow(ws, [12, 16, 7] + [12] * len(metr_keys))

ws = sheet("Per-run raw (TCGA)")
tk = ["accuracy", "precision_macro", "recall_macro", "f1_macro", "auc_macro_ovr",
      "train_time", "epochs_run"]
put_row(ws, 1, ["Method", "Fold"] + tk, HDR, FILL)
r = 2
for key in sorted(TC):
    meth, fold = key.split("|")
    v = TC[key]
    if "error" in v:
        continue
    put_row(ws, r, [meth, fold] + [round(v[mk], 5) if mk in v else "" for mk in tk])
    r += 1
autow(ws, [16, 7] + [13] * len(tk))

# ---- Datasets ----
ws = sheet("Datasets")
st = agg["dataset_stats"]
put_row(ws, 1, ["Dataset", "Type", "Task", "Train", "Valid", "Test", "Features",
                "Labels", "Cardinality", "Density"], HDR, FILL)
r = 2
for d in DATASETS:
    s = st[d]
    put_row(ws, r, [d, "Yeast gene expression (FunCat)" if d != "yeast14"
                    else "Yeast gene expression (Mulan/Elisseeff-Weston)", "Multi-label",
                    s["train"], s["valid"], s["test"], s["features"], s["labels"],
                    round(s["cardinality"], 3), round(s["density"], 4)])
    r += 1
t = st["tcga"]
put_row(ws, r, ["TCGA-PANCAN", "RNA-seq (human tumors)", "Multiclass",
                "801 (5-fold CV)", "", "", t["genes"], 5, 1.0, 0.2])
r += 2
ws.cell(r, 1, "TCGA class distribution: " + ", ".join(f"{k}: {v}" for k, v in t["classes"].items())).font = TXT
autow(ws, [16, 28, 14, 10, 10, 10, 10, 9, 12, 10])

wb.save("results/AHGSL_results.xlsx")
print("workbook saved: results/AHGSL_results.xlsx")
